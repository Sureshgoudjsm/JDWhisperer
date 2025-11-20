import os
import json
import csv
import io
import re
import shutil
import pandas as pd
import datetime
from dotenv import load_dotenv
import google.generativeai as genai
import streamlit as st

# --- New Libraries for File Parsing ---
try:
    from pypdf import PdfReader
    from docx import Document
except ImportError:
    st.error("Please install missing libraries: pip install pypdf python-docx")
    st.stop()

# Optional Google Sheets libs
try:
    import gspread
    from google.oauth2.service_account import Credentials
    HAS_GSHEETS_LIBS = True
except ImportError:
    HAS_GSHEETS_LIBS = False

# ------------------------------
#  PAGE CONFIG (brand)
# ------------------------------
st.set_page_config(
    page_title="JD Whisperer",
    layout="wide",
    page_icon="🤫"
)

# ------------------------------
#  LOGO HANDLING
# ------------------------------
ORIG_UPLOADED_PATH = '/mnt/data/A_logo_in_digital_vector_art_format_for_"JD_Whispe.png'
CLEAN_LOCAL_LOGO = 'jd_whisperer_logo.png'
LOCAL_LOGO_PATH = CLEAN_LOCAL_LOGO
# Fallback URL
DRIVE_FILE_ID = "1DJoP8qI8X5mgFnuB3eQueC_WbX7_AT5n"
LOGO_URL = f"https://drive.google.com/uc?export=view&id={DRIVE_FILE_ID}"

# ------------------------------
#  ENV + AI CONFIG
# ------------------------------
load_dotenv()
API_KEY = os.getenv("GOOGLE_API_KEY") or st.secrets.get("GOOGLE_API_KEY", None)
if not API_KEY:
    st.error("CRITICAL: GOOGLE_API_KEY not set. Add to env or Streamlit Secrets.")
    st.stop()

genai.configure(api_key=API_KEY)

@st.cache_resource
def get_model():
    return genai.GenerativeModel('gemini-2.5-flash')

# ------------------------------
#  FILE PARSING
# ------------------------------
def extract_text_from_pdf(file):
    try:
        reader = PdfReader(file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    except Exception as e:
        return f"Error reading PDF: {e}"

def extract_text_from_docx(file):
    try:
        doc = Document(file)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        return f"Error reading DOCX: {e}"

# ------------------------------
#  GOOGLE SHEETS
# ------------------------------
FIELD_ORDER = [
    "date_contacted","hr_name","phone_number","email_id","role_position",
    "recruiter_company","client_company","location","job_type","mode_of_contact",
    "interview_mode","interview_scheduled_date","round_1_details","round_2_details",
    "ctc_offered_expected","status","next_follow_up_date","review_notes",
    "extracted_keywords","match_score","skill_gap_analysis","prep_hint"
]

def _get_gsheets_creds_and_id():
    try:
        sheet_id = st.secrets.get("GOOGLE_SHEET_ID", None)
        sa_info = st.secrets.get("GOOGLE_SERVICE_ACCOUNT", None)
        if not sheet_id or not sa_info:
            return None, None
        if isinstance(sa_info, str):
            creds_dict = json.loads(sa_info)
        else:
            creds_dict = dict(sa_info)
        return creds_dict, sheet_id
    except Exception:
        return None, None

@st.cache_resource
def get_gsheets_worksheet():
    if not HAS_GSHEETS_LIBS:
        return None
    creds_dict, sheet_id = _get_gsheets_creds_and_id()
    if not creds_dict or not sheet_id:
        return None
    scopes = ["https://www.googleapis.com/auth/spreadsheets"]
    creds = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    client = gspread.authorize(creds)
    sh = client.open_by_key(sheet_id)
    ws = sh.sheet1
    try:
        existing_values = ws.get_all_values()
        if not existing_values:
            headers = ["timestamp_utc"] + FIELD_ORDER
            ws.append_row(headers, value_input_option="USER_ENTERED")
    except Exception:
        pass
    return ws

def save_history_to_gsheets(result: dict):
    try:
        ws = get_gsheets_worksheet()
        if ws is None:
            return
        timestamp = datetime.datetime.utcnow().isoformat()
        row = [timestamp] + [result.get(k, "") for k in FIELD_ORDER]
        ws.append_row(row, value_input_option="USER_ENTERED")
    except Exception as e:
        st.warning(f"Could not save to Google Sheets: {e}")

def load_history_dataframe() -> pd.DataFrame:
    df = None
    try:
        ws = get_gsheets_worksheet()
    except Exception:
        ws = None
    if ws is not None:
        try:
            records = ws.get_all_records()
            if records:
                df = pd.DataFrame(records)
        except Exception as e:
            st.warning(f"Could not load history from Google Sheets: {e}")
    if (df is None or df.empty) and st.session_state.get("history"):
        df = pd.DataFrame(st.session_state.get("history"))
    if df is None:
        df = pd.DataFrame()
    return df

# ------------------------------
#  SESSION STATE & RESET LOGIC
# ------------------------------
def init_state():
    if 'app_state' not in st.session_state:
        st.session_state.app_state = {
            'current_view': 'start',
            'profile_data': "",
            'job_description': "",
            'skills_data': "",
            'analysis_result': None,
            'generated_email': "" # New feature state
        }
    if 'history' not in st.session_state:
        st.session_state['history'] = []
    if 'mode' not in st.session_state:
        st.session_state['mode'] = 'Analyze'

def soft_reset():
    """Keeps Profile, clears JD and Results."""
    st.session_state.app_state['current_view'] = 'map'
    # Note: We deliberately DO NOT clear 'profile_data'
    st.session_state.app_state['job_description'] = ""
    st.session_state.app_state['skills_data'] = ""
    st.session_state.app_state['analysis_result'] = None
    st.session_state.app_state['generated_email'] = ""

def hard_reset():
    """Clears EVERYTHING."""
    st.session_state.app_state = {
        'current_view': 'start',
        'profile_data': "",
        'job_description': "",
        'skills_data': "",
        'analysis_result': None,
        'generated_email': ""
    }

init_state()

# ------------------------------
#  AI PROMPTS
# ------------------------------
EXTRACTION_PROMPT = """
You are an expert data extraction assistant for job seekers. Your task is to analyze the provided texts:
1) Job Details (JD, email, call notes) and
2) Applicant Skills (Resume/Summary).

You MUST:
- Infer as many fields as possible from context.
- Use "Not specified" if you truly cannot infer a value.

CRITICAL: Return a single, valid JSON object only.

FIELD-SPECIFIC RULES:
- "interview_scheduled_date" as "YYYY-MM-DD"
- "skill_gap_analysis": VERY SHORT (2-3 sentences max)
- "prep_hint": 1-2 short sentences

JSON Keys (all must be present):
""" + ", ".join(FIELD_ORDER + ["match_score","skill_gap_analysis","prep_hint"]) + """

Input:
***
{text_input}
***
Return ONLY the JSON object.
"""

EMAIL_PROMPT = """
You are a top-tier career coach. Write a short, punchy, professional "Cold Email" or "Cover Letter" for the candidate to send to the recruiter.

Context:
- Candidate Profile: {profile}
- Job Role: {role}
- Company: {company}
- Recruiter Name: {recruiter} (If unknown, use "Hiring Manager")
- Key Matching Skills: {keywords}

Instructions:
- Tone: Professional, confident, but not arrogant.
- Length: Under 150 words.
- Structure: Hook -> Value Prop -> Call to Action.
- Do not include placeholders like [Your Name] if you can infer it from the profile.

Output: Just the email body text.
"""

# ------------------------------
#  HELPERS
# ------------------------------
def safe_json_from_response(text: str) -> dict:
    cleaned = text.strip().replace('```json', '').replace('```', '')
    match = re.search(r'\{.*\}', cleaned, re.DOTALL)
    if not match:
        return json.loads(cleaned)
    return json.loads(match.group(0))

def keep_first_sentences(text: str, max_sentences: int = 3) -> str:
    if not isinstance(text, str):
        return text
    sentences = re.split(r'(?<=[.!?])\s+', text.strip())
    sentences = [s for s in sentences if s]
    if not sentences:
        return text
    return " ".join(sentences[:max_sentences])

def process_recruiter_text(text_to_process: str) -> dict:
    model = get_model()
    prompt_with_input = EXTRACTION_PROMPT.format(text_input=text_to_process)
    try:
        response = model.generate_content(prompt_with_input)
        raw_text = response.text or ""
        parsed = safe_json_from_response(raw_text)
        for key in ("skill_gap_analysis", "prep_hint"):
            if key in parsed and isinstance(parsed[key], str):
                parsed[key] = keep_first_sentences(parsed[key], 3)
        return parsed
    except json.JSONDecodeError:
        return {"error": f"Invalid JSON returned. Raw: {response.text if 'response' in locals() else 'N/A'}"}
    except Exception as e:
        return {"error": f"Error: {e}"}

def generate_email_logic():
    res = st.session_state.app_state['analysis_result']
    profile = st.session_state.app_state['profile_data']
    
    if not res or not profile: return "Error: Missing data."

    prompt = EMAIL_PROMPT.format(
        profile=profile[:3000], # Limit context
        role=res.get('role_position', 'the role'),
        company=res.get('client_company', 'your company'),
        recruiter=res.get('hr_name', 'Hiring Manager'),
        keywords=res.get('extracted_keywords', 'my skills')
    )
    
    model = get_model()
    try:
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        return f"Could not generate email: {e}"

def create_ics_file(details: dict) -> str:
    date_str = details.get("interview_scheduled_date")
    if not date_str or date_str == "Not specified":
        return ""
    try:
        start_date = datetime.datetime.strptime(date_str, "%Y-%m-%d").replace(hour=10, minute=0)
        end_date = start_date + datetime.timedelta(hours=1)
        dt_format = "%Y%m%dT%H%M%S"
        summary = f"Interview: {details.get('role_position', 'Job')} @ {details.get('client_company', 'Client')}"
        description = f"Role: {details.get('role_position','N/A')}\\nCompany: {details.get('client_company','N/A')}\\nNotes: {details.get('review_notes','')}"
        ics_content = (
            "BEGIN:VCALENDAR\nVERSION:2.0\nPRODID:-//JD Whisperer//EN\nBEGIN:VEVENT\n"
            f"UID:{datetime.datetime.now().strftime(dt_format)}-{hash(summary)}\n"
            f"DTSTAMP:{datetime.datetime.now().strftime(dt_format)}\n"
            f"DTSTART:{start_date.strftime(dt_format)}\n"
            f"DTEND:{end_date.strftime(dt_format)}\n"
            f"SUMMARY:{summary}\nDESCRIPTION:{description}\nEND:VEVENT\nEND:VCALENDAR"
        )
        return ics_content
    except Exception:
        return ""

def sanitize_df_for_streamlit(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for col in df.columns:
        if df[col].map(lambda x: isinstance(x, (dict, list, set, tuple))).any():
            df[col] = df[col].map(lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list, set, tuple)) else x)
    return df

# ------------------------------
#  BRAND CSS
# ------------------------------
# ------------------------------
#  BRAND CSS
# ------------------------------
def load_brand_css():
    # NOTE: We removed the 'f' before the string so Python doesn't try to interpret the CSS
    st.markdown("""
    <link href="https://fonts.googleapis.com/css2?family=Inter:wght@300;400;600&family=Manrope:wght@400;600;800&display=swap" rel="stylesheet">
    <style>
    :root {
        --bg: #0B1E37;
        --surface: #122642;
        --muted: #9EACBE;
        --text: #F5F9FF;
        --accent-1: #4C8CFF;
        --accent-2: #6A5CFF;
        --accent-3: #00D4D0;
    }
    body {
        font-family: 'Inter', 'Manrope', sans-serif;
        background: var(--bg) !important;
        color: var(--text) !important;
    }
    .stApp .block-container {
        padding: 2rem 2rem 3rem 2rem !important;
        background: linear-gradient(180deg, rgba(11,30,55,0.95) 0%, rgba(17,38,66,0.95) 100%);
        border-radius: 8px;
    }
    .main-header {
        text-align: left;
        padding: 1rem 0;
        display: flex;
        align-items: center;
        gap: 1rem;
    }
    .main-header img.logo {
        height: 56px;
    }
    .main-header h1 {
        margin: 0;
        font-family: 'Manrope', sans-serif;
        font-weight: 800;
        font-size: 34px;
        color: var(--text);
        letter-spacing: -0.02em;
        background: linear-gradient(90deg, var(--accent-1), var(--accent-2), var(--accent-3));
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .main-header p {
        margin: 0;
        color: var(--muted);
        font-size: 14px;
    }
    .mind-map-card {
        background: linear-gradient(180deg, rgba(22,38,60,0.6), rgba(17,30,45,0.45));
        border: 1px solid rgba(76,140,255,0.12);
        border-radius: 12px;
        padding: 18px;
        text-align: center;
        min-height: 250px;
    }
    .stButton > button {
        background: linear-gradient(90deg, var(--accent-1), var(--accent-2)) !important;
        color: white !important;
        border-radius: 8px !important;
        padding: 10px 14px !important;
        font-weight: 700 !important;
        border: none !important;
    }
    .stButton > button:hover {
        filter: brightness(1.03);
    }
    .stMetric > div > div > div {
        color: var(--accent-3) !important;
    }
    .whisper-divider {
        height: 6px;
        border-radius: 6px;
        background: linear-gradient(90deg, var(--accent-1), var(--accent-2), var(--accent-3));
        margin: 16px 0;
    }
    /* Start Screen specific centering */
    .start-container {
        display: flex;
        flex-direction: column;
        align-items: center;
        justify-content: center;
        min-height: 60vh;
        text-align: center;
        gap: 1rem;
    }
    </style>
    """, unsafe_allow_html=True)

# ------------------------------
#  UI VIEWS
# ------------------------------
def try_show_logo(width=72):
    if os.path.exists(LOCAL_LOGO_PATH):
        try:
            st.image(LOCAL_LOGO_PATH, width=width)
            return
        except Exception:
            pass
    try:
        st.image(LOGO_URL, width=width)
    except Exception:
        pass

def draw_start_view():
    st.markdown('<div class="main-header">', unsafe_allow_html=True)
    try_show_logo(width=50)
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="start-container">', unsafe_allow_html=True)
    st.markdown('<h1 style="font-size: 4rem; margin-bottom: 0;">JD Whisperer</h1>', unsafe_allow_html=True)
    st.markdown('<p style="font-size: 1.2rem; color: #9EACBE; margin-bottom: 2rem;">Intelligently Map Your Next Career Move</p>', unsafe_allow_html=True)
    
    _, btn_col, _ = st.columns([5, 2, 5])
    with btn_col:
        if st.button("🚀 Start Mapping", use_container_width=True):
            st.session_state.app_state['current_view'] = 'map'
            st.rerun()
            
    st.markdown('</div>', unsafe_allow_html=True)

def draw_map_view():
    st.markdown('<div style="display:flex;gap:1rem;align-items:center;margin-bottom:1rem">', unsafe_allow_html=True)
    st.markdown('<h2 style="margin:0;color:var(--text)">Build Your Career Mind Map</h2>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    col1, col2, col3 = st.columns(3, gap="medium")
    
    # --- CARD 1: PROFILE (With Persistence & File Upload) ---
    with col1:
        st.markdown('<div class="mind-map-card"><h3>👤 Your Profile</h3><p>Upload Resume or Paste Text</p></div>', unsafe_allow_html=True)
        
        # Show status if data is present
        if st.session_state.app_state['profile_data']:
             st.success("✅ Profile Loaded")
             if st.button("❌ Clear Profile", key="clear_profile"):
                 st.session_state.app_state['profile_data'] = ""
                 st.rerun()
        
        tab_up, tab_txt = st.tabs(["📂 Upload", "✍️ Paste"])
        
        with tab_up:
            uploaded_file = st.file_uploader("Upload PDF/DOCX", type=['pdf', 'docx'], key="resume_file")
            if uploaded_file is not None:
                with st.spinner("Reading file..."):
                    if uploaded_file.name.endswith('.pdf'):
                        text = extract_text_from_pdf(uploaded_file)
                    else:
                        text = extract_text_from_docx(uploaded_file)
                    
                    if text:
                        st.session_state.app_state['profile_data'] = text
                        st.success("Resume parsed successfully!")
                        # We don't rerun immediately to let user see success message, 
                        # but value is updated in state.

        with tab_txt:
            val = st.text_area("Paste Text", 
                             value=st.session_state.app_state['profile_data'], 
                             height=150, 
                             key="profile_input_area", 
                             label_visibility="collapsed", 
                             placeholder="e.g., 5+ years in Python, AWS...")
            if val != st.session_state.app_state['profile_data']:
                st.session_state.app_state['profile_data'] = val

    # --- CARD 2: JD (Clears on Soft Reset) ---
    with col2:
        st.markdown('<div class="mind-map-card"><h3>📄 Job Description</h3><p>Paste the JD, recruiter email, or call notes.</p></div>', unsafe_allow_html=True)
        st.markdown("<br>", unsafe_allow_html=True)
        jd_text = st.text_area("Job Description", 
                             value=st.session_state.app_state['job_description'],
                             height=230, 
                             key="jd_input", 
                             label_visibility="collapsed", 
                             placeholder="Paste JD here...")
        st.session_state.app_state['job_description'] = jd_text

    # --- CARD 3: NOTES (Clears on Soft Reset) ---
    with col3:
        st.markdown('<div class="mind-map-card"><h3>📝 Extra Notes</h3><p>CTC, notice period, call summary.</p></div>', unsafe_allow_html=True)
        st.markdown("<br>", unsafe_allow_html=True)
        skills_text = st.text_area("Skill Assessment", 
                                 value=st.session_state.app_state['skills_data'],
                                 height=230, 
                                 key="skills_input", 
                                 label_visibility="collapsed", 
                                 placeholder="Additional notes...")
        st.session_state.app_state['skills_data'] = skills_text

    st.markdown('<div class="whisper-divider"></div>', unsafe_allow_html=True)

    has_profile = bool(st.session_state.app_state['profile_data'].strip())
    has_jd = bool(st.session_state.app_state['job_description'].strip())
    
    if not (has_profile and has_jd):
        st.info("Please provide **Your Profile** and **Job Description** to generate analysis.")

    if st.button("✨ Generate Analysis", disabled=not (has_profile and has_jd), use_container_width=True):
        combined_text = (
            f"--- APPLICANT SKILLS ---\n{st.session_state.app_state['profile_data']}\n\n"
            f"--- JOB DETAILS ---\n{st.session_state.app_state['job_description']}\n\n"
            f"--- ADDITIONAL NOTES ---\n{st.session_state.app_state['skills_data']}"
        )
        with st.spinner("🧠 JD Whisperer is analyzing..."):
            result = process_recruiter_text(combined_text)
            st.session_state.app_state['analysis_result'] = result
            if "error" not in result:
                st.session_state['history'].append(result)
                save_history_to_gsheets(result)
            st.session_state.app_state['current_view'] = 'results'
            st.rerun()

def draw_results_view():
    st.markdown('<div style="display:flex;justify-content:space-between;align-items:center">', unsafe_allow_html=True)
    left, right = st.columns([1,4])
    with left:
        try_show_logo(width=64)
    with right:
        st.markdown('<h2 style="margin:0">Analysis Results</h2>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)

    result = st.session_state.app_state.get('analysis_result')
    if not result:
        st.error("No analysis found.")
        if st.button("Back"): soft_reset(); st.rerun()
        return
    if "error" in result:
        st.error(result.get("error"))
        if st.button("Back"): soft_reset(); st.rerun()
        return

    raw_score = result.get('match_score','N/A')
    try:
        match_score_display = f"{int(str(raw_score).replace('%','').strip())}%"
    except Exception:
        match_score_display = str(raw_score)

    # TABS for Results View
    tab_dash, tab_email, tab_data = st.tabs(["📊 Dashboard", "✉️ Email Generator", "💾 Raw Data"])

    with tab_dash:
        col1, col2 = st.columns([1,2], gap="large")
        with col1:
            st.metric(label="Match Score", value=match_score_display)
            st.info(f"**Gap:** {result.get('skill_gap_analysis','-')}")
            st.success(f"**Hint:** {result.get('prep_hint','-')}")
        with col2:
            st.subheader("Summary")
            st.write(f"**Role:** {result.get('role_position')}")
            st.write(f"**Company:** {result.get('client_company')}")
            st.write(f"**Keywords:** {result.get('extracted_keywords')}")

    with tab_email:
        st.markdown("### ✉️ Instant Cold Email Draft")
        st.caption("Based on your profile and this specific job description.")
        
        if st.button("✨ Draft Email to HR", key="gen_email_btn"):
            with st.spinner("Writing email..."):
                email_text = generate_email_logic()
                st.session_state.app_state['generated_email'] = email_text
        
        if st.session_state.app_state['generated_email']:
            st.text_area("Copy this:", value=st.session_state.app_state['generated_email'], height=300)
    
    with tab_data:
        df_display = pd.DataFrame([result]).T
        df_display.columns = ["Value"]
        df_display = sanitize_df_for_streamlit(df_display)
        st.dataframe(df_display, use_container_width=True)
        
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=result.keys())
        writer.writeheader()
        writer.writerow(result)
        csv_data = output.getvalue()
        st.download_button("📄 Download CSV", data=csv_data, file_name="job_details.csv", mime="text/csv")

    st.markdown("---")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("⬅️ Edit Inputs (Keep Data)"):
            st.session_state.app_state['current_view'] = 'map'
            st.rerun()
    with c2:
        # Soft Reset: Clears JD/Notes but KEEPS Profile
        if st.button("🔄 Analyze Next Job (Keep Profile)"):
            soft_reset()
            st.rerun()

def draw_history_view():
    st.markdown("## 📊 History Dashboard")
    df = load_history_dataframe()
    if df.empty:
        st.info("No history found yet.")
        return
    st.dataframe(df, use_container_width=True, hide_index=True)

# ------------------------------
#  MAIN ROUTER
# ------------------------------
def main():
    load_brand_css()

    # Sidebar
    with st.sidebar:
        st.markdown("<div style='display:flex;align-items:center;gap:8px'>", unsafe_allow_html=True)
        try_show_logo(width=48)
        st.markdown("<div><strong>JD Whisperer</strong><br><small style='color:var(--muted)'>Decode any job description</small></div>", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)

        st.markdown("#### Mode")
        col_a, col_b = st.columns(2)
        with col_a:
            if st.button("Analyze", use_container_width=True, key="mode_analyze_btn"):
                st.session_state['mode'] = "Analyze"
                st.rerun()
        with col_b:
            if st.button("History", use_container_width=True, key="mode_history_btn"):
                st.session_state['mode'] = "History"
                st.rerun()
        
        st.markdown("---")
        if st.button("🗑️ Hard Reset (Clear All)"):
            hard_reset()
            st.rerun()
        
        creds_dict, sheet_id = _get_gsheets_creds_and_id()
        if HAS_GSHEETS_LIBS and creds_dict and sheet_id:
            st.caption("✅ GSheets Connected")

    mode = st.session_state.get('mode', 'Analyze')
    if mode == "History":
        draw_history_view()
        return

    view = st.session_state.app_state.get('current_view','start')
    if view == 'start':
        draw_start_view()
    elif view == 'map':
        draw_map_view()
    elif view == 'results':
        draw_results_view()
    else:
        reset_app_state()
        draw_start_view()

if __name__ == "__main__":
    main()